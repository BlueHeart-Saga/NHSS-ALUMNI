import asyncio
import csv
import io
from fastapi import APIRouter, Depends, HTTPException, Response
from app.core.database import get_db, build_school_filter
from app.core.cache import ttl_cache
from app.schemas.models import DashboardReportResponse
from app.middleware.auth import require_roles

from bson import ObjectId
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

router = APIRouter(prefix="/reports", tags=["Reports & Analytics"])

@router.get("/summary", response_model=DashboardReportResponse)
async def get_dashboard_summary(
    current_user: dict = Depends(require_roles(["SCHOOL_ADMIN", "BATCH_COORDINATOR", "SUPER_ADMIN", "DEVELOPER", "PLATFORM_DEVELOPER"]))
):
    school_id = current_user.get("school_id") or "ALL"
    cache_key = f"reports:summary:{school_id}"
    cached = ttl_cache.get(cache_key)
    if cached is not None:
        return cached

    db = get_db()
    base_filter = await build_school_filter(school_id)

    # 1. Single aggregation pipeline on db.alumni to aggregate status counts without session explosion
    alumni_pipeline = []
    if base_filter:
        alumni_pipeline.append({"$match": base_filter})
    alumni_pipeline.append({
        "$group": {
            "_id": "$verification_status",
            "count": {"$sum": 1}
        }
    })

    batch_query = {**base_filter}
    batch_query["$or"] = [{"status": "ACTIVE"}, {"status": {"$exists": False}}, {"status": None}]

    events_query = dict(base_filter)
    checkins_query = dict(base_filter)

    # Run aggregation and count queries concurrently in parallel
    alumni_counts, active_batches, upcoming_events, recent_checkins = await asyncio.gather(
        db.alumni.aggregate(alumni_pipeline).to_list(length=100),
        db.batches.count_documents(batch_query),
        db.events.count_documents(events_query),
        db.checkins.count_documents(checkins_query)
    )

    total_alumni = 0
    verified_alumni = 0
    pending_alumni = 0

    for item in alumni_counts:
        status_val = item.get("_id")
        cnt = item.get("count", 0)
        total_alumni += cnt
        if status_val in ["APPROVED", "VERIFIED"]:
            verified_alumni += cnt
        elif status_val == "PENDING":
            pending_alumni += cnt

    turnout_pct = (recent_checkins / verified_alumni * 100) if verified_alumni > 0 else (100.0 if recent_checkins > 0 else 0.0)

    res = DashboardReportResponse(
        total_alumni=total_alumni,
        verified_alumni=verified_alumni,
        pending_alumni=pending_alumni,
        active_batches=active_batches,
        upcoming_events=upcoming_events,
        recent_checkins_count=recent_checkins,
        attendance_turnout_percentage=round(turnout_pct, 1)
    )

    ttl_cache.set(cache_key, res, ttl=30)
    return res

@router.get("/export-alumni")
async def export_alumni_csv(
    current_user: dict = Depends(require_roles(["SCHOOL_ADMIN", "SUPER_ADMIN", "DEVELOPER", "PLATFORM_DEVELOPER"]))
):
    db = get_db()
    school_id = current_user.get("school_id")
    filter_query = await build_school_filter(school_id)

    cursor = db.alumni.find(filter_query).sort("passing_year", -1)
    alumni_list = await cursor.to_list(length=10000)

    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")

    # Complete 44-column Header matching all editable spreadsheet & profile fields
    writer.writerow([
        "Alumni ID",
        "Full Name",
        "Name in Tamil",
        "Mobile Number",
        "Country Code",
        "Gender",
        "Date of Birth",
        "Email",
        "Blood Group",
        "Father Name",
        "Mother Name",
        "Current City",
        "Current State",
        "Address",
        "Country",
        "School Name",
        "Joining Year",
        "Passing Year",
        "Leaving Class",
        "Admission Number",
        "Roll No",
        "Section",
        "No Higher Ed",
        "College Name",
        "Degree",
        "Custom Degree",
        "Department",
        "College Reg No",
        "College Joining Yr",
        "College Passing Yr",
        "Employment Status",
        "Company Name",
        "Designation",
        "Industry",
        "Total Experience",
        "Skills",
        "LinkedIn URL",
        "Instagram URL",
        "WhatsApp Number",
        "Website URL",
        "Profile Photo URL",
        "Is Volunteer",
        "Willing to Donate",
        "Verification Status"
    ])

    def excel_safe_text(value) -> str:
        """Wrap string values in an Excel ="..." formula so Excel treats them as TEXT.
        This prevents long numeric/hex strings (like MongoDB ObjectIds and mobile
        numbers with + country code) from being converted into scientific notation.
        Internal double quotes are escaped by doubling them.
        """
        if value is None:
            return ""
        s = str(value)
        if s == "":
            return ""
        escaped = s.replace('"', '""')
        return f'="{escaped}"'

    for a in alumni_list:
        raw_id = str(a.get("_id", ""))
        raw_mobile = a.get("mobile", "") or ""
        raw_country_code = a.get("country_code", "") or "91"
        raw_adm_no = a.get("admission_number", "") or ""
        raw_roll_no = a.get("roll_no", "") or ""
        raw_college_reg = a.get("college_register_no") or a.get("register_number") or ""
        raw_whatsapp = a.get("whatsapp_number", "") or ""

        # Skills list or string
        skills_val = a.get("skills", "")
        if isinstance(skills_val, list):
            skills_str = ", ".join(str(s) for s in skills_val if s)
        else:
            skills_str = str(skills_val) if skills_val is not None else ""

        # Higher education flag
        no_high_ed = a.get("no_higher_education")
        if isinstance(no_high_ed, bool):
            no_high_ed_str = "YES" if no_high_ed else "NO"
        else:
            no_high_ed_str = str(no_high_ed).strip().upper() if no_high_ed else "NO"

        writer.writerow([
            excel_safe_text(raw_id),
            a.get("full_name", ""),
            a.get("name_ta") or a.get("full_name_ta", ""),
            excel_safe_text(raw_mobile) if raw_mobile else "",
            str(raw_country_code),
            a.get("gender", ""),
            a.get("date_of_birth") or a.get("dob", ""),
            a.get("email", ""),
            a.get("blood_group", ""),
            a.get("father_name", ""),
            a.get("mother_name", ""),
            a.get("current_city", ""),
            a.get("current_state") or a.get("state", ""),
            a.get("address", ""),
            a.get("country", "India"),
            a.get("school_name", ""),
            a.get("joining_year") or a.get("admission_year") or "",
            a.get("passing_year", ""),
            a.get("leaving_class", ""),
            excel_safe_text(raw_adm_no) if raw_adm_no else "",
            excel_safe_text(raw_roll_no) if raw_roll_no else "",
            a.get("section", ""),
            no_high_ed_str,
            a.get("college_name") or a.get("institution_name", ""),
            a.get("degree", ""),
            a.get("custom_degree") or a.get("other_degree", ""),
            a.get("department") or a.get("stream", ""),
            excel_safe_text(raw_college_reg) if raw_college_reg else "",
            a.get("college_joining_year") or "",
            a.get("college_passing_year") or "",
            a.get("employment_status", ""),
            a.get("company_name") or a.get("company", ""),
            a.get("profession") or a.get("designation") or a.get("position", ""),
            a.get("industry", ""),
            a.get("total_experience") or a.get("experience_years", ""),
            skills_str,
            a.get("linkedin_url", ""),
            a.get("instagram_url", ""),
            excel_safe_text(raw_whatsapp) if raw_whatsapp else "",
            a.get("website_url", ""),
            a.get("profile_photo_url", ""),
            a.get("is_volunteer", "NO"),
            a.get("willing_to_donate", "NO"),
            a.get("verification_status", "APPROVED")
        ])

    # Prepend UTF-8 BOM so Microsoft Excel detects UTF-8 and renders Tamil correctly
    csv_data = "\ufeff" + output.getvalue()

    return Response(
        content=csv_data.encode("utf-8"),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=alumni_roster.csv"}
    )

@router.get("/export-alumni-excel")
async def export_alumni_excel(
    current_user: dict = Depends(require_roles(["SCHOOL_ADMIN", "SUPER_ADMIN", "DEVELOPER", "PLATFORM_DEVELOPER"]))
):
    db = get_db()
    school_id = current_user.get("school_id")
    filter_query = await build_school_filter(school_id)

    cursor = db.alumni.find(filter_query).sort("passing_year", -1)
    alumni_list = await cursor.to_list(length=10000)

    wb = Workbook()
    ws = wb.active
    ws.title = "Alumni Directory"
    ws.views.sheetView[0].showGridLines = True

    headers = [
        "Alumni ID",
        "Full Name",
        "Name in Tamil",
        "Mobile Number",
        "Country Code",
        "Gender",
        "Date of Birth",
        "Email",
        "Blood Group",
        "Father Name",
        "Mother Name",
        "Current City",
        "Current State",
        "Address",
        "Country",
        "School Name",
        "Joining Year",
        "Passing Year",
        "Leaving Class",
        "Admission Number",
        "Roll No",
        "Section",
        "No Higher Ed",
        "College Name",
        "Degree",
        "Custom Degree",
        "Department",
        "College Reg No",
        "College Joining Yr",
        "College Passing Yr",
        "Employment Status",
        "Company Name",
        "Designation",
        "Industry",
        "Total Experience",
        "Skills",
        "LinkedIn URL",
        "Instagram URL",
        "WhatsApp Number",
        "Website URL",
        "Profile Photo URL",
        "Is Volunteer",
        "Willing to Donate",
        "Verification Status"
    ]

    # Styles
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")  # Slate-800 Navy
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border_side = Side(border_style="thin", color="E2E8F0")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    font_regular = Font(name="Calibri", size=10)
    font_bold_approved = Font(name="Calibri", size=10, bold=True, color="166534")
    font_bold_pending = Font(name="Calibri", size=10, bold=True, color="854D0E")
    font_bold_rejected = Font(name="Calibri", size=10, bold=True, color="991B1B")

    align_left = Alignment(horizontal="left", vertical="center")
    align_center = Alignment(horizontal="center", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # Header Row
    ws.append(headers)
    ws.row_dimensions[1].height = 28

    for col_idx in range(1, len(headers) + 1):
        h_cell = ws.cell(row=1, column=col_idx)
        h_cell.font = header_font
        h_cell.fill = header_fill
        h_cell.alignment = header_align
        h_cell.border = cell_border

    # Center-aligned column indices (1-indexed)
    center_cols = {5, 6, 7, 9, 17, 18, 19, 20, 21, 22, 23, 28, 29, 30, 42, 43, 44}
    # Text-formatted column indices (prevent scientific notation and leading zero stripping)
    text_cols = {1, 4, 15, 20, 21, 28, 39}
    # Integer/Numeric columns
    numeric_cols = {17, 18, 29, 30}

    # Data Rows
    for r_idx, a in enumerate(alumni_list, start=2):
        raw_id = str(a.get("_id", ""))
        raw_mobile = str(a.get("mobile", "") or "")
        raw_country_code = str(a.get("country_code", "") or "91")
        raw_adm_no = str(a.get("admission_number", "") or "")
        raw_roll_no = str(a.get("roll_no", "") or "")
        raw_college_reg = str(a.get("college_register_no") or a.get("register_number") or "")
        raw_whatsapp = str(a.get("whatsapp_number", "") or "")

        skills_val = a.get("skills", "")
        if isinstance(skills_val, list):
            skills_str = ", ".join(str(s) for s in skills_val if s)
        else:
            skills_str = str(skills_val) if skills_val is not None else ""

        no_high_ed = a.get("no_higher_education")
        if isinstance(no_high_ed, bool):
            no_high_ed_str = "YES" if no_high_ed else "NO"
        else:
            no_high_ed_str = str(no_high_ed).strip().upper() if no_high_ed else "NO"

        def to_year_val(val):
            if val is None or val == "":
                return ""
            try:
                return int(float(val))
            except (ValueError, TypeError):
                return str(val)

        joining_yr = to_year_val(a.get("joining_year") or a.get("admission_year"))
        passing_yr = to_year_val(a.get("passing_year"))
        col_join_yr = to_year_val(a.get("college_joining_year"))
        col_pass_yr = to_year_val(a.get("college_passing_year"))

        status = str(a.get("verification_status", "APPROVED")).upper()

        row_data = [
            raw_id,
            str(a.get("full_name", "") or ""),
            str(a.get("name_ta") or a.get("full_name_ta", "") or ""),
            raw_mobile,
            raw_country_code,
            str(a.get("gender", "") or ""),
            str(a.get("date_of_birth") or a.get("dob", "") or ""),
            str(a.get("email", "") or ""),
            str(a.get("blood_group", "") or ""),
            str(a.get("father_name", "") or ""),
            str(a.get("mother_name", "") or ""),
            str(a.get("current_city", "") or ""),
            str(a.get("current_state") or a.get("state", "") or ""),
            str(a.get("address", "") or ""),
            str(a.get("country", "India") or "India"),
            str(a.get("school_name", "") or ""),
            joining_yr,
            passing_yr,
            str(a.get("leaving_class", "") or ""),
            raw_adm_no,
            raw_roll_no,
            str(a.get("section", "") or ""),
            no_high_ed_str,
            str(a.get("college_name") or a.get("institution_name", "") or ""),
            str(a.get("degree", "") or ""),
            str(a.get("custom_degree") or a.get("other_degree", "") or ""),
            str(a.get("department") or a.get("stream", "") or ""),
            raw_college_reg,
            col_join_yr,
            col_pass_yr,
            str(a.get("employment_status", "") or ""),
            str(a.get("company_name") or a.get("company", "") or ""),
            str(a.get("profession") or a.get("designation") or a.get("position", "") or ""),
            str(a.get("industry", "") or ""),
            str(a.get("total_experience") or a.get("experience_years", "") or ""),
            skills_str,
            str(a.get("linkedin_url", "") or ""),
            str(a.get("instagram_url", "") or ""),
            raw_whatsapp,
            str(a.get("website_url", "") or ""),
            str(a.get("profile_photo_url", "") or ""),
            str(a.get("is_volunteer", "NO") or "NO"),
            str(a.get("willing_to_donate", "NO") or "NO"),
            status
        ]

        ws.append(row_data)
        ws.row_dimensions[r_idx].height = 20

        for col_idx in range(1, len(row_data) + 1):
            cell = ws.cell(row=r_idx, column=col_idx)
            cell.border = cell_border
            cell.font = font_regular

            # Text format to avoid scientific notation
            if col_idx in text_cols:
                cell.number_format = '@'
                cell.alignment = align_left if col_idx not in center_cols else align_center
            elif col_idx in numeric_cols and isinstance(cell.value, int):
                cell.number_format = '0'
                cell.alignment = align_center
            elif col_idx in center_cols:
                cell.alignment = align_center
            else:
                cell.alignment = align_left

            # Status column special font styling
            if col_idx == 44:
                if status == "APPROVED":
                    cell.font = font_bold_approved
                elif status == "PENDING":
                    cell.font = font_bold_pending
                elif status in ("REJECTED", "SUSPENDED"):
                    cell.font = font_bold_rejected

    # Freeze header row
    ws.freeze_panes = "A2"

    # AutoFilter across all headers and data rows
    max_row = max(len(alumni_list) + 1, 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max_row}"

    # Auto-adjust column widths based on content + header length (capped 12 to 40)
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=alumni_roster.xlsx"}
    )


@router.get("/export-attendance/{event_id}")
async def export_event_attendance_csv(
    event_id: str,
    current_user: dict = Depends(require_roles(["SCHOOL_ADMIN", "BATCH_COORDINATOR"]))
):
    db = get_db()
    school_id = current_user["school_id"]
    from bson import ObjectId

    cursor = db.event_attendance.find({"event_id": event_id})
    att_list = await cursor.to_list(length=5000)

    output = io.StringIO()
    writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")

    writer.writerow([
        "Alumni Name", "Batch Year", "Admission Number", "RSVP Status",
        "Adult Guests", "Child Guests", "Total Guests Expected", "Check-in Status", "Check-in Time"
    ])

    for att in att_list:
        alumni = await db.alumni.find_one({"_id": ObjectId(att["alumni_id"])})
        checkin = await db.checkins.find_one({"event_id": event_id, "alumni_id": att["alumni_id"]})

        writer.writerow([
            alumni.get("full_name", "Unknown") if alumni else "Unknown",
            alumni.get("passing_year", "") if alumni else "",
            alumni.get("admission_number", "") if alumni else "",
            att.get("rsvp_status", ""),
            att.get("adults_count", 1),
            att.get("children_count", 0),
            att.get("total_guests", 1),
            "Checked In" if checkin else "Not Checked In",
            checkin["checked_in_at"].strftime("%Y-%m-%d %I:%M %p") if checkin else ""
        ])

    # Prepend UTF-8 BOM for Excel compatibility
    csv_data = "\ufeff" + output.getvalue()

    return Response(
        content=csv_data.encode("utf-8"),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename=event_{event_id}_attendance.csv"}
    )